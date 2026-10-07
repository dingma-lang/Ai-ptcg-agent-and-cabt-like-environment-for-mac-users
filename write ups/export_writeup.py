"""Export the approved Markdown article to DOCX and PDF, including its figure.

Uses python-docx, reportlab, markdown-it-py and matplotlib already installed
in the local document toolchain. The source article is never rewritten.
"""
from __future__ import annotations

import argparse
import html
import os
import re
from pathlib import Path

os.environ.setdefault('MPLCONFIGDIR', '/private/tmp/ptcg-writeup-matplotlib')
import matplotlib
matplotlib.use('Agg')
from matplotlib import mathtext
from matplotlib.font_manager import FontProperties
from markdown_it import MarkdownIt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Table, TableStyle, KeepTogether, Spacer


MD = MarkdownIt('commonmark').enable('table')
FONT_DIR = Path('/System/Library/Fonts/Supplemental')


def parse_article(source):
    body, separator, notes = source.partition('## Evidence and References')
    if not separator:
        raise ValueError('Expected the v8 reference section')
    references = re.findall(r'^\[\^([^\]]+)\]:\s*(.+)$', notes, re.MULTILINE)
    used = set(re.findall(r'\[\^([^\]]+)\]', body))
    if used != {key for key, _ in references}:
        raise ValueError('Missing or unused references')
    body = re.sub(r'\[\^([^\]]+)\]', r'[[\1]](#ref-\1)', body)
    tokens = MD.parse(body)
    blocks = []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type in ('heading_open', 'paragraph_open'):
            content = tokens[index + 1]
            kind = 'title' if token.tag == 'h1' else 'heading' if token.tag == 'h2' else 'paragraph'
            if content.content.startswith('$$') and content.content.endswith('$$'):
                blocks.append(('equation', content.content[2:-2]))
            elif content.children and content.children[0].type == 'image':
                blocks.append(('image', content.children[0].attrGet('src')))
            else:
                if content.content.startswith('*Figure 1.'):
                    kind = 'caption'
                elif len(blocks) == 1 and kind == 'paragraph':
                    kind = 'subtitle'
                blocks.append((kind, content.children))
            index += 3
        elif token.type == 'table_open':
            rows = []
            index += 1
            while tokens[index].type != 'table_close':
                if tokens[index].type == 'tr_open':
                    rows.append([])
                elif tokens[index].type == 'inline':
                    rows[-1].append(tokens[index].children)
                index += 1
            blocks.append(('table', rows))
            index += 1
        else:
            raise ValueError('Unsupported article block: ' + token.type)
    blocks.append(('heading', MD.parseInline('Evidence and References')[0].children))
    for key, value in references:
        blocks.append(('reference', (key, MD.parseInline(value)[0].children)))
    return blocks


def pdf_inline(tokens):
    result = []
    for token in tokens:
        if token.type in ('text', 'code_inline'):
            result.append(html.escape(token.content))
        elif token.type in ('softbreak', 'hardbreak'):
            result.append(' ')
        elif token.type in ('em_open', 'em_close', 'strong_open', 'strong_close'):
            result.append({'em_open':'<i>', 'em_close':'</i>', 'strong_open':'<b>', 'strong_close':'</b>'}[token.type])
        elif token.type == 'link_open':
            target = html.escape(token.attrGet('href'), quote=True)
            result.append('<link href="' + target + '" color="#236B65">')
        elif token.type == 'link_close':
            result.append('</link>')
        else:
            raise ValueError('Unsupported inline token: ' + token.type)
    return ''.join(result)


def word_inline(paragraph, tokens):
    bold = italic = False
    link = None
    for token in tokens:
        if token.type == 'strong_open': bold = True
        elif token.type == 'strong_close': bold = False
        elif token.type == 'em_open': italic = True
        elif token.type == 'em_close': italic = False
        elif token.type == 'link_open': link = token.attrGet('href')
        elif token.type == 'link_close': link = None
        elif token.type in ('text', 'code_inline', 'softbreak', 'hardbreak'):
            value = ' ' if token.type in ('softbreak', 'hardbreak') else token.content
            if link:
                hyperlink = OxmlElement('w:hyperlink')
                if link.startswith('#'):
                    hyperlink.set(qn('w:anchor'), link[1:])
                else:
                    hyperlink.set(qn('r:id'), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                properties = OxmlElement('w:rPr')
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '236B65')
                properties.append(color)
                if bold: properties.append(OxmlElement('w:b'))
                if italic: properties.append(OxmlElement('w:i'))
                run.append(properties)
                text = OxmlElement('w:t')
                text.text = value
                text.set(qn('xml:space'), 'preserve')
                run.append(text)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            else:
                run = paragraph.add_run(value)
                run.bold, run.italic = bold, italic
        else:
            raise ValueError('Unsupported Word inline token: ' + token.type)


def bookmark(paragraph, name, identifier):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(identifier))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def prepare_equations(blocks, assets):
    assets.mkdir(parents=True, exist_ok=True)
    rendered = {}
    for kind, data in blocks:
        if kind != 'equation': continue
        expression = data.replace(r'\text{no Basic in seven}', r'\mathrm{no\ Basic\ in\ seven}')
        target = assets / ('equation_' + str(len(rendered) + 1) + '.png')
        mathtext.math_to_image('$' + expression + '$', str(target), dpi=240,
                              prop=FontProperties(size=14), color='#243A3D')
        rendered[data] = target
    return rendered


def word_export(blocks, source, equations, output):
    doc = Document()
    doc.core_properties.title = 'From Legal Choices to Coherent Turns'
    doc.core_properties.subject = 'Cross-deck strategy write-up, version 8'
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Georgia', Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for style, size in [('Title',25), ('Subtitle',12), ('Heading 1',14), ('Caption',9)]:
        doc.styles[style].font.name = 'Arial'
        doc.styles[style].font.size = Pt(size)
        doc.styles[style].font.color.rgb = RGBColor.from_string('236B65' if style == 'Heading 1' else '243A3D')
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Cross-deck strategy | v8   ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for index, (kind, data) in enumerate(blocks):
        if kind == 'table':
            table = doc.add_table(rows=0, cols=len(data[0]))
            table.autofit = False
            for column, width in zip(table.columns, [1.45, 2.35, 3.2]): column.width = Inches(width)
            table.style = 'Light Shading Accent 1'
            for row_index, row in enumerate(data):
                cells = table.add_row().cells
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement('w:cantSplit'))
                if row_index == 0: props.append(OxmlElement('w:tblHeader'))
                for cell, content in zip(cells, row):
                    word_inline(cell.paragraphs[0], content)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(5)
                    cell.paragraphs[0].paragraph_format.space_before = Pt(5)
                    for run in cell.paragraphs[0].runs:
                        run.font.name, run.font.size = 'Arial', Pt(9)
                        run.bold = row_index == 0
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), 'EAF1F0' if row_index == 0 else 'FFFFFF')
                    cell._tc.get_or_add_tcPr().append(shading)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif kind in ('image', 'equation'):
            path = equations[data] if kind == 'equation' else source.parent / data
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if kind == 'image':
                width = 7.0
                paragraph.paragraph_format.keep_with_next = True
            else:
                pixels = ImageReader(str(path)).getSize()[0]
                width = min(6.7, pixels / 240)
            paragraph.add_run().add_picture(str(path), width=Inches(width))
        elif kind == 'reference':
            key, tokens = data
            paragraph = doc.add_paragraph()
            bookmark(paragraph, 'ref-' + key, index)
            paragraph.add_run('[' + key + '] ').bold = True
            word_inline(paragraph, tokens)
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs: run.font.size = Pt(9)
        else:
            style = {'title':'Title', 'subtitle':'Subtitle', 'heading':'Heading 1', 'caption':'Caption'}.get(kind, 'Normal')
            paragraph = doc.add_paragraph(style=style)
            word_inline(paragraph, data)
            paragraph.paragraph_format.widow_control = True
            if kind in ('heading', 'title', 'subtitle'):
                paragraph.paragraph_format.keep_with_next = True
    doc.save(output)


def pdf_export(blocks, source, equations, output):
    for family in ['Arial', 'Georgia']:
        for suffix, filename in [('',family), ('-Bold',family+' Bold'), ('-Italic',family+' Italic'), ('-BoldItalic',family+' Bold Italic')]:
            pdfmetrics.registerFont(TTFont(family+suffix, str(FONT_DIR / (filename+'.ttf'))))
        pdfmetrics.registerFontFamily(family, normal=family, bold=family+'-Bold', italic=family+'-Italic', boldItalic=family+'-BoldItalic')
    base = ParagraphStyle('Body', fontName='Georgia', fontSize=10.5, leading=14.4,
                          textColor=colors.HexColor('#243A3D'), spaceAfter=8, allowWidows=0, allowOrphans=0)
    styles = {
        'paragraph':base,
        'title':ParagraphStyle('Title',parent=base,fontName='Arial-Bold',fontSize=25,leading=29,spaceAfter=9,keepWithNext=True),
        'subtitle':ParagraphStyle('Subtitle',parent=base,fontName='Arial-Italic',fontSize=11.5,leading=15,spaceAfter=14,keepWithNext=True),
        'heading':ParagraphStyle('Heading',parent=base,fontName='Arial-Bold',fontSize=14,leading=18,spaceBefore=14,spaceAfter=9,textColor=colors.HexColor('#236B65'),keepWithNext=True),
        'caption':ParagraphStyle('Caption',parent=base,fontName='Arial-Italic',fontSize=8.5,leading=11,spaceAfter=12),
        'reference':ParagraphStyle('Reference',parent=base,fontName='Arial',fontSize=9,leading=12,spaceAfter=8),
        'cell':ParagraphStyle('Cell',parent=base,fontName='Arial',fontSize=9,leading=12,spaceAfter=0),
    }
    story = []
    skip_caption = False
    for index, (kind, data) in enumerate(blocks):
        if kind == 'caption' and skip_caption:
            skip_caption = False
            continue
        if kind == 'table':
            rows = [[Paragraph(pdf_inline(cell), styles['cell']) for cell in row] for row in data]
            table = Table(rows, colWidths=[104,168,232], repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF1F0')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),
                ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
                ('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),
                ('LINEBELOW',(0,0),(-1,-1),.5,colors.HexColor('#CDDCD8')),
            ]))
            story.extend([table, Spacer(1,12)])
        elif kind in ('image', 'equation'):
            path = equations[data] if kind == 'equation' else source.parent / data
            width, height = ImageReader(str(path)).getSize()
            draw_width = min(484, width * 72 / 240) if kind == 'equation' else 504
            image = Image(str(path), width=draw_width, height=height*draw_width/width)
            image.hAlign = 'CENTER'
            group = [Spacer(1,5), image, Spacer(1,8)]
            if kind == 'image' and index+1 < len(blocks) and blocks[index+1][0] == 'caption':
                group.append(Paragraph(pdf_inline(blocks[index+1][1]),styles['caption']))
                skip_caption = True
            story.append(KeepTogether(group))
        elif kind == 'reference':
            key, content = data
            story.append(Paragraph('<a name="ref-'+key+'"/><b>['+key+']</b> '+pdf_inline(content), styles['reference']))
        else:
            story.append(Paragraph(pdf_inline(data), styles[kind]))
    document = SimpleDocTemplate(str(output), pagesize=(612,792), leftMargin=54,rightMargin=54,
                                 topMargin=48,bottomMargin=48, title='From Legal Choices to Coherent Turns')
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#D5E1DD'))
        canvas.line(54,35,558,35)
        canvas.setFont('Arial',8)
        canvas.setFillColor(colors.HexColor('#66756F'))
        canvas.drawString(54,23,'Cross-deck strategy | v8')
        canvas.drawRightString(558,23,str(doc.page))
        canvas.restoreState()
    document.build(story,onFirstPage=footer,onLaterPages=footer)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=Path(__file__).with_name('cross_deck_strategy_draft_v8.md'))
    args = parser.parse_args()
    source = args.source.resolve()
    blocks = parse_article(source.read_text(encoding='utf-8'))
    equations = prepare_equations(blocks, source.parent / 'exports' / 'v8_assets')
    word_export(blocks, source, equations, source.with_suffix('.docx'))
    pdf_export(blocks, source, equations, source.with_suffix('.pdf'))
    print('Created:', source.with_suffix('.docx'))
    print('Created:', source.with_suffix('.pdf'))
    print('Embedded figures:', sum(kind == 'image' for kind,_ in blocks))
    print('Equations:',len(equations))
    print('References:',sum(kind == 'reference' for kind,_ in blocks))


if __name__ == '__main__':
    main()
