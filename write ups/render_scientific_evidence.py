"""Render figures and a source-linked workbook from the measured analysis outputs."""

import csv
import json
import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/ptcg-writeup-matplotlib")
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties


BASE = Path(__file__).resolve().parent / "scientific_data"


def add_sheet(workbook, name, rows):
    sheet = workbook.create_sheet(name)
    if not rows:
        return sheet
    sheet.append(list(rows[0]))
    for row in rows:
        sheet.append([json.dumps(value) if isinstance(value, (list, dict)) else value for value in row.values()])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343A40")
    for column in sheet.columns:
        width = min(65, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return sheet


def csv_rows(path):
    with path.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    integers = {"card_id", "count", "hp", "decision", "turn", "player", "instance", "water",
                "lightning", "total_energy", "typed_deficit", "options", "selected", "selected_type", "top"}
    decimals = {"selected_score", "top_score", "score_gap", "rank_loss"}
    for row in rows:
        for key, value in row.items():
            if key in integers and value:
                row[key] = int(value)
            elif key in decimals and value:
                row[key] = float(value)
            elif value in ("True", "False"):
                row[key] = value == "True"
    return rows


def main():
    summary = json.loads((BASE / "scientific_summary.json").read_text())
    probes = json.loads((BASE / "archive_probes.json").read_text())
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11})
    opening = summary["opening_probabilities"][2:]
    labels = ["Finizen", "Finizen + Dratini", "Finizen + Water", "Finizen + Dratini + Water",
              "Buddy-Buddy Poffin", "Two or more Palafin ex"]
    fig, ax = plt.subplots(figsize=(10, 5.1))
    fig.subplots_adjust(left=0.34, right=0.95, bottom=0.24, top=0.84)
    values = [100 * item["probability"] for item in opening]
    ax.barh(range(len(labels)), values, color="#287c77", height=0.58)
    ax.set_yticks(range(len(labels)), labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 80)
    ax.set_xlabel("Exact probability in a seven-card hand containing a Basic (%)")
    ax.grid(axis="x", color="#e6e6e6", linewidth=0.7)
    ax.set_axisbelow(True)
    for index, value in enumerate(values):
        ax.text(value + 1, index, "{:.2f}%".format(value), va="center", fontsize=10)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    fig.suptitle("Opening access from the actual 60-card deck", fontsize=16, x=0.05, ha="left", weight="bold")
    fig.text(0.05, 0.11, "7 Basics: 4 Finizen, 2 Dratini, 1 Tatsugiri. 6 Water, 4 Poffin, 4 Palafin ex.", fontsize=10)
    fig.text(0.05, 0.055, "Uniform shuffle; conditional on a valid opening. No bonus draws, searches, or turn-one draw included.", fontsize=9)
    fig.savefig(BASE / "opening_access.png", dpi=180, facecolor="white")
    plt.close(fig)

    configurations = probes["configurations"]
    order = [("1.1", "full"), ("1.2", "full"), ("1.3", "without_before_commit"), ("1.3", "full")]
    ordered = [next(c for c in configurations if (c["version"], c["mode"]) == pair) for pair in order]
    fig, ax = plt.subplots(figsize=(10, 5.6))
    fig.subplots_adjust(left=0.09, right=0.96, bottom=0.28, top=0.84)
    for offset, field, color, label in [(-0.19, "expert_matches", "#287c77", "Original expert index"),
                                        (0.19, "regression_matches", "#667078", "Agreed regression target")]:
        ax.bar([i + offset for i in range(4)], [r[field] for r in ordered], width=0.34, color=color, label=label)
        for i, row in enumerate(ordered):
            ax.text(i + offset, row[field] + 0.25, str(row[field]), ha="center", fontsize=11)
    ax.set_xticks(range(4), ["1.1", "1.2", "1.3\npre-check disabled", "1.3\nfull"])
    ax.set_ylim(0, 16)
    ax.set_yticks(range(0, 15, 2))
    ax.set_ylabel("Matched decisions / 14 reviewed windows")
    ax.grid(axis="y", color="#e6e6e6", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, loc="upper left")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.suptitle("Preparation-check ablation on fixed observations", fontsize=16, x=0.06, ha="left", weight="bold")
    fig.text(0.06, 0.11, "Disabling only before_commit_choice changes six decisions. Each input is repeated three times.", fontsize=10)
    fig.text(0.06, 0.065, "Selected development cases, not held-out accuracy or win rate. Decision 136 has two distinct labels.", fontsize=9)
    fig.savefig(BASE / "preparation_ablation.png", dpi=180, facecolor="white")
    plt.close(fig)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    add_sheet(workbook, "Read Me", [
        {"topic": "Scope", "detail": "Historical observations, exact probabilities, and isolated archive probes. No new games."},
        {"topic": "Experimental units", "detail": "29 unique inputs x 4 configurations x 3 repeats = 348 calls, not 348 independent cases."},
        {"topic": "Ablation", "detail": "Only before_commit_choice returns None in a temporary extracted 1.3 process. Source ZIP unchanged."},
        {"topic": "Limitations", "detail": "Reviewed development set; cold match memory; no alternate trajectories; repeated card sightings are correlated."},
        {"topic": "Model versus data", "detail": "Opening probabilities use exact deck combinatorics. Opening Observations contains 20 selected recorded hands."},
        {"topic": "Formula", "detail": "Probabilities: precomputed numeric result plus an Excel numerator/denominator formula; cached formula values require spreadsheet recalculation."},
    ])
    sheet = add_sheet(workbook, "Probabilities", summary["opening_probabilities"])
    sheet.cell(1, 6, "Excel verification")
    sheet.cell(1, 6).font = Font(bold=True, color="FFFFFF")
    sheet.cell(1, 6).fill = PatternFill("solid", fgColor="343A40")
    for row in range(2, len(opening) + 4):
        sheet.cell(row, 6, "=C{0}/D{0}".format(row))
        sheet.cell(row, 2).number_format = "0.00%"
        sheet.cell(row, 6).number_format = "0.00%"
    sheet.column_dimensions["F"].width = 22
    add_sheet(workbook, "Trace Summary", summary["trace_aggregates"])
    sheet = add_sheet(workbook, "Cross Deck Openings", summary["cross_deck_openings"])
    for row in range(2, 6):
        sheet.cell(row, 6).number_format = "0.00%"
        sheet.cell(row, 7).number_format = "0.00%"
    add_sheet(workbook, "Probe Summary", configurations)
    add_sheet(workbook, "Probe Calls", probes["records"])
    add_sheet(workbook, "Energy Summary", summary["energy_aggregates"])
    for name, filename in [("Deck Counts", "deck_counts.csv"), ("Opening Observations", "opening_observations.csv"),
                           ("Score Windows", "score_discordance.csv"), ("Attacker Snapshots", "attacker_snapshots.csv")]:
        add_sheet(workbook, name, csv_rows(BASE / filename))
    add_sheet(workbook, "Battle15 Transitions", [{"from_to_pass": key, "count": value}
                                               for key, value in summary["battle15_transitions"].items()])
    sources = [
        {"source": "SciPy", "reference": "Hypergeometric distribution", "location": "https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.hypergeom.html", "use": "Sampling without replacement"},
        {"source": "NIST", "reference": "Confidence intervals, 7.2.4.1", "location": "https://www.itl.nist.gov/div898/handbook/prc/section2/prc241.htm", "use": "Previous Wilson intervals"},
        {"source": "Jiang and Li (2016)", "reference": "Doubly Robust Off-policy Value Evaluation", "location": "https://proceedings.mlr.press/v48/jiang16.html", "use": "Scope limits of fixed-observation evaluation"},
        {"source": "Local analysis", "reference": "scientific_analysis.py", "location": "../scientific_analysis.py", "use": "All combinatorics, trace, opening, and Energy measurements"},
        {"source": "Local experiment", "reference": "probe_archived_policies.py", "location": "../probe_archived_policies.py", "use": "Fixed-observation version comparison and ablation"},
    ]
    add_sheet(workbook, "Sources", sources)
    add_sheet(workbook, "Input Hashes", summary["source_manifest"])
    add_sheet(workbook, "Probe Input Hashes", probes["sources"])
    workbook.save(BASE / "scientific_evidence.xlsx")
    print("Wrote opening_access.png, preparation_ablation.png, scientific_evidence.xlsx")


if __name__ == "__main__":
    main()
